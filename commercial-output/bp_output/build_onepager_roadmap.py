"""
LAIQ Product Roadmap One-Pager
Single-slide roadmap in the same BP template style.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

BASE = os.path.dirname(os.path.abspath(__file__))
PARENT = os.path.dirname(BASE)
LOGO = os.path.join(PARENT, "laiq_logo.png")
OUT = os.path.join(BASE, "LAIQ_Product_Roadmap_OnePager.pptx")

NAVY     = RGBColor(0x0A, 0x16, 0x28)
NAVY2    = RGBColor(0x12, 0x24, 0x40)
TEAL     = RGBColor(0x07, 0x5F, 0x5C)
TEAL_LT  = RGBColor(0xDC, 0xEF, 0xE9)
TEAL_MID = RGBColor(0x8A, 0xBB, 0xB6)
TEAL_DIM = RGBColor(0x4A, 0x80, 0x7C)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xF4, 0xF7, 0xF6)
LGRAY    = RGBColor(0xE2, 0xE8, 0xE6)
MID      = RGBColor(0x52, 0x62, 0x5E)
PALE     = RGBColor(0xA8, 0xB8, 0xB4)
AMBER    = RGBColor(0xE6, 0x7E, 0x22)
DARK     = RGBColor(0x17, 0x20, 0x1E)

I = Inches
prs = Presentation()
prs.slide_width = I(13.33)
prs.slide_height = I(7.5)
BLANK = prs.slide_layouts[6]
FONT = "Arial"


def rect(s, l, t, w, h, fill=None, line=None, lw=0.6):
    sh = s.shapes.add_shape(1, I(l), I(t), I(w), I(h))
    sh.fill.solid() if fill else sh.fill.background()
    if fill:
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def ellipse(s, l, t, w, h, fill=None, line=None, lw=0.6):
    sh = s.shapes.add_shape(9, I(l), I(t), I(w), I(h))
    sh.fill.solid() if fill else sh.fill.background()
    if fill:
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def txt(s, text, l, t, w, h, size=11, bold=False, color=DARK,
        align=PP_ALIGN.LEFT, wrap=True, italic=False):
    tb = s.shapes.add_textbox(I(l), I(t), I(w), I(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return tb


def bg(s, color=WHITE):
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = color


def add_logo(s, rgt=13.13, bot=7.44, h=0.26):
    if not os.path.exists(LOGO):
        return
    w = h * (930 / 430)
    s.shapes.add_picture(LOGO, I(rgt - w), I(bot - h), width=I(w), height=I(h))


def footer(s, label=""):
    rect(s, 0, 7.18, 13.33, 0.04, TEAL_MID)
    if label:
        txt(s, label, 0.18, 7.23, 6, 0.21, size=7.5, color=PALE)
    txt(s, "LAIQ  ·  Confidential", 7, 7.23, 6.13, 0.21,
        size=7.5, color=PALE, align=PP_ALIGN.RIGHT)
    add_logo(s)


def left_bar(s, color=TEAL, accent=TEAL_LT):
    rect(s, 0, 0, 0.32, 7.5, color)
    rect(s, 0.32, 0, 0.07, 7.5, accent)


def slide_kicker(s, kicker, title, tcolor=NAVY, kcolor=TEAL):
    txt(s, kicker, 0.55, 0.20, 12.6, 0.26, size=8.5, bold=True, color=kcolor)
    txt(s, title, 0.55, 0.44, 12.6, 0.76, size=24, bold=True, color=tcolor)
    rect(s, 0.55, 1.20, 12.60, 0.04, TEAL_MID)


s = prs.slides.add_slide(BLANK)
bg(s)
left_bar(s, TEAL)
slide_kicker(s, "PRODUCT ROADMAP",
             "From Oil & Gas Inspection Copilot to Industrial Service Network", TEAL)

txt(s,
    "Start with oil & gas tank inspection, expand into wider industrial inspection, then build multi-agent maintenance intelligence and a service network connecting clients, inspectors, workshops, OEMs, and spare-part providers.",
    0.55, 1.32, 12.2, 0.42, size=10.2, color=MID)

rect(s, 0.55, 1.78, 12.62, 4.98, OFFWHITE, TEAL_MID, 0.5)
txt(s, "Four-stage roadmap", 5.38, 1.96, 2.4, 0.22, size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

stages = [
    (
        "1",
        "Oil & Gas\nInspection Copilot",
        "Leave site with a report-ready tank inspection package",
        "Prove the wedge",
        TEAL,
        [
            "Offline tank capture",
            "UT, checklist, findings, photos",
            "Location + MFL/NDT reference",
            "AI report draft",
        ],
    ),
    (
        "2",
        "Industrial\nInspection Copilot",
        "Apply the same inspection-to-report workflow across more assets and industries",
        "Expand the workflow",
        TEAL_DIM,
        [
            "Vessels, piping, valves, pumps",
            "Utilities, power, marine, heavy equipment",
            "Configurable inspection templates",
            "More reports, more devices, more customers",
        ],
    ),
    (
        "3",
        "Multi-Agent Maintenance\n& Reliability Intelligence",
        "Turn inspection and maintenance history into weak-point, lifetime, FMEA, and repair-priority decisions",
        "Decide what to fix next",
        AMBER,
        [
            "Report, review, planning, FMEA agents",
            "Reliability + weak-point analysis",
            "CMMS / work-order linkage",
            "Premium intelligence module",
        ],
    ),
    (
        "4",
        "Inspection-to-\nService Network",
        "Connect findings to workshops, OEMs, spare parts, service execution, and follow-up",
        "Act through the network",
        NAVY2,
        [
            "Repair package + service routing",
            "Quote / work-order flow",
            "Spare-part demand + OEM feedback",
            "Partner / marketplace revenue",
        ],
    ),
]

lefts = [0.82, 3.92, 7.02, 10.12]
for i, (num, title, body, tag, accent, bullets) in enumerate(stages):
    l = lefts[i]
    rect(s, l, 2.22, 2.58, 3.18, WHITE, TEAL_MID, 0.5)
    rect(s, l, 2.22, 2.58, 0.06, accent)
    ellipse(s, l + 0.14, 2.38, 0.40, 0.40, accent)
    txt(s, num, l + 0.14, 2.38, 0.40, 0.40, size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, title, l + 0.60, 2.34, 1.82, 0.42, size=10.3, bold=True, color=DARK)
    txt(s, body, l + 0.16, 2.92, 2.18, 0.64, size=8.5, color=MID, wrap=True)
    rect(s, l + 0.16, 3.70, 2.10, 0.24, accent)
    txt(s, tag, l + 0.22, 3.75, 1.98, 0.14, size=7.4, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    by = 4.10
    for b in bullets:
        txt(s, f"•  {b}", l + 0.16, by, 2.18, 0.18, size=7.5, color=DARK)
        by += 0.20
    if i < 3:
        txt(s, "▶", l + 2.66, 3.36, 0.22, 0.20, size=20, bold=True, color=TEAL_MID, align=PP_ALIGN.CENTER)

rect(s, 0.90, 5.58, 12.00, 0.04, LGRAY)
txt(s, "Structured inspection data spine", 0.90, 5.72, 2.7, 0.18, size=9.8, bold=True, color=NAVY)

chips = [
    ("Field data", TEAL),
    ("Report package", TEAL_DIM),
    ("Asset history", AMBER),
    ("Reliability decisions", NAVY2),
    ("Service execution", NAVY),
]
cx = 3.02
for label, accent in chips:
    w = len(label) * 0.08 + 0.34
    rect(s, cx, 5.70, w, 0.28, TEAL_LT, accent, 0.5)
    txt(s, label, cx + 0.10, 5.76, w - 0.20, 0.14, size=8.0, bold=True, color=accent, align=PP_ALIGN.CENTER)
    cx += w + 0.10

rect(s, 0.90, 6.18, 12.02, 0.56, NAVY)
txt(s,
    "Roadmap principle: keep one structured inspection data spine, then expand from oil & gas reports to industrial inspection, maintenance intelligence, reliability management, and service execution.",
    1.08, 6.34, 11.66, 0.18, size=10, bold=True, color=TEAL_LT, align=PP_ALIGN.CENTER)

footer(s, "Tank Inspection · One-page product roadmap")

prs.save(OUT)
print(f"✅  Saved → {OUT}")
