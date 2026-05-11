"""
LAIQ Tank Inspection Solution One-Pager
Single-slide companion to the pain points page.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

BASE = os.path.dirname(os.path.abspath(__file__))
PARENT = os.path.dirname(BASE)
LOGO = os.path.join(PARENT, "laiq_logo.png")
OUT = os.path.join(BASE, "LAIQ_Tank_Solution_OnePager.pptx")

NAVY     = RGBColor(0x0A, 0x16, 0x28)
TEAL     = RGBColor(0x07, 0x5F, 0x5C)
TEAL_LT  = RGBColor(0xDC, 0xEF, 0xE9)
TEAL_MID = RGBColor(0x8A, 0xBB, 0xB6)
TEAL_DIM = RGBColor(0x4A, 0x80, 0x7C)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xF4, 0xF7, 0xF6)
LGRAY    = RGBColor(0xE2, 0xE8, 0xE6)
MID      = RGBColor(0x52, 0x62, 0x5E)
PALE     = RGBColor(0xA8, 0xB8, 0xB4)
AMBER    = RGBColor(0xE6, 0x7E, 0x22)
DARK     = RGBColor(0x17, 0x20, 0x1E)

I = Inches
prs = Presentation()
prs.slide_width = I(13.33)
prs.slide_height = I(7.5)
BLANK = prs.slide_layouts[6]
FONT = "Arial"


def rect(s, l, t, w, h, fill=None, line=None, lw=0.6):
    sh = s.shapes.add_shape(1, I(l), I(t), I(w), I(h))
    sh.fill.solid() if fill else sh.fill.background()
    if fill:
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def ellipse(s, l, t, w, h, fill=None, line=None, lw=0.6):
    sh = s.shapes.add_shape(9, I(l), I(t), I(w), I(h))
    sh.fill.solid() if fill else sh.fill.background()
    if fill:
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def txt(s, text, l, t, w, h, size=11, bold=False, color=DARK,
        align=PP_ALIGN.LEFT, wrap=True, italic=False):
    tb = s.shapes.add_textbox(I(l), I(t), I(w), I(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return tb


def bg(s, color=WHITE):
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = color


def add_logo(s, rgt=13.13, bot=7.44, h=0.26):
    if not os.path.exists(LOGO):
        return
    w = h * (930 / 430)
    s.shapes.add_picture(LOGO, I(rgt - w), I(bot - h), width=I(w), height=I(h))


def footer(s, label=""):
    rect(s, 0, 7.18, 13.33, 0.04, TEAL_MID)
    if label:
        txt(s, label, 0.18, 7.23, 6, 0.21, size=7.5, color=PALE)
    txt(s, "LAIQ  ·  Confidential", 7, 7.23, 6.13, 0.21,
        size=7.5, color=PALE, align=PP_ALIGN.RIGHT)
    add_logo(s)


def left_bar(s, color=TEAL, accent=TEAL_LT):
    rect(s, 0, 0, 0.32, 7.5, color)
    rect(s, 0.32, 0, 0.07, 7.5, accent)


def slide_kicker(s, kicker, title, tcolor=NAVY, kcolor=TEAL):
    txt(s, kicker, 0.55, 0.20, 12.6, 0.26, size=8.5, bold=True, color=kcolor)
    txt(s, title, 0.55, 0.44, 12.6, 0.76, size=24, bold=True, color=tcolor)
    rect(s, 0.55, 1.20, 12.60, 0.04, TEAL_MID)


def bullet_list(s, items, l, t, w, gap=0.30, size=9.2, color=DARK):
    for i, item in enumerate(items):
        txt(s, f"•  {item}", l, t + i * gap, w, 0.20, size=size, color=color)


s = prs.slides.add_slide(BLANK)
bg(s)
left_bar(s, TEAL)
slide_kicker(s, "LAIQ SOLUTION",
             "One Structured Field Workflow Drives the Final Report Package", TEAL)

txt(s,
    "One guided field workflow captures every surface. A canonical inspection record then drives report assembly, evidence linking, and structured output.",
    0.55, 1.32, 12.2, 0.34, size=10.5, color=MID)

# Left: tank / surfaces / capture model
rect(s, 0.55, 1.74, 4.70, 4.98, OFFWHITE, TEAL_MID, 0.5)
txt(s, "Inspection surfaces captured in one model", 0.82, 1.92, 2.9, 0.22,
    size=10, bold=True, color=NAVY)

tx, ty, tw, th = 1.22, 2.20, 2.12, 3.52
ellipse(s, tx, ty, tw, 0.44, TEAL_LT, TEAL_MID, 0.8)
rect(s, tx, ty + 0.22, tw, th - 0.44, TEAL_DIM, TEAL_MID, 0.8)
ellipse(s, tx, ty + th - 0.44, tw, 0.44, TEAL_LT, TEAL_MID, 0.8)
for i in range(1, 6):
    y = ty + 0.22 + i * (th - 0.88) / 6
    rect(s, tx, y, tw, 0.015, TEAL_LT)

txt(s, "Roof", 0.72, 2.34, 0.40, 0.18, size=9.5, bold=True, color=DARK, align=PP_ALIGN.RIGHT)
txt(s, "Shell", 0.70, 3.06, 0.42, 0.18, size=9.5, bold=True, color=DARK, align=PP_ALIGN.RIGHT)
txt(s, "Nozzles", 0.58, 4.08, 0.54, 0.18, size=9.5, bold=True, color=DARK, align=PP_ALIGN.RIGHT)
txt(s, "Bottom /\nMFL", 0.52, 5.06, 0.60, 0.34, size=9.5, bold=True, color=DARK, align=PP_ALIGN.RIGHT)

callouts = [
    ("Roof", "UT plates from roof surfaces\nwith ring / section or plate logic", 3.62, 2.18),
    ("Shell", "UT by strake + N/E/S/W\nwith optional precise mapping", 3.62, 3.02),
    ("Nozzles", "Shell and roof nozzle UT\ncaptured in structured tables", 3.62, 4.04),
    ("Bottom / MFL", "Import MFL metadata and\nreference into the same job", 3.62, 5.10),
]
for _, body, l, t in callouts:
    rect(s, l, t, 1.34, 0.64, WHITE, TEAL_MID, 0.4)
    txt(s, body, l + 0.10, t + 0.10, 1.14, 0.40, size=7.8, color=DARK)
    txt(s, "•", l - 0.16, t + 0.18, 0.12, 0.12, size=16, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

# Center arrows
txt(s, "▶", 5.38, 2.84, 0.34, 0.22, size=24, bold=True, color=TEAL_MID, align=PP_ALIGN.CENTER)
txt(s, "▶", 5.38, 4.20, 0.34, 0.22, size=24, bold=True, color=TEAL_MID, align=PP_ALIGN.CENTER)

# Right: three-step LAIQ solve flow
card_left = 5.86
card_w = 2.30
gap = 0.24
tops = 1.86

for idx, (title, subtitle, bullets, fill, body_fill) in enumerate([
    ("1  Field Capture", "Guided on-site workflow", [
        "Tank & surface configuration",
        "UT rows: shell, roof, nozzle",
        "In-context finding capture",
        "Photo evidence linked at source",
        "Coarse or precise location",
    ], TEAL, WHITE),
    ("2  Inspection Record", "Canonical structured data", [
        "Asset hierarchy: tank, surface, plate",
        "Location: strake, bearing, X/Y",
        "Findings: type, severity, evidence",
        "MFL metadata and attachment",
        "Completeness and API checks",
    ], TEAL_DIM, WHITE),
    ("3  Report Package", "Automated output assembly", [
        "Shell / roof UT tables (auto-filled)",
        "Findings register + linked photos",
        "Shell plate layout map",
        "MFL section with reference",
        "JSON + CSV export",
    ], AMBER, WHITE),
]):
    l = card_left + idx * (card_w + gap)
    rect(s, l, tops, card_w, 0.54, fill)
    txt(s, title, l + 0.12, tops + 0.10, card_w - 0.24, 0.18,
        size=10.8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, subtitle, l + 0.12, tops + 0.30, card_w - 0.24, 0.14,
        size=7.8, color=TEAL_LT if fill != AMBER else WHITE, align=PP_ALIGN.CENTER, italic=True)
    rect(s, l, tops + 0.60, card_w, 2.36, body_fill, TEAL_MID, 0.4)
    bullet_list(s, bullets, l + 0.12, tops + 0.78, card_w - 0.24, gap=0.32, size=8.5, color=DARK)

txt(s, "▶", 8.24, 2.96, 0.22, 0.20, size=18, bold=True, color=TEAL_MID, align=PP_ALIGN.CENTER)
txt(s, "▶", 10.78, 2.96, 0.22, 0.20, size=18, bold=True, color=TEAL_MID, align=PP_ALIGN.CENTER)

rect(s, 5.86, 5.18, 7.08, 1.02, TEAL_LT, TEAL_MID, 0.5)
txt(s, "What LAIQ changes", 6.08, 5.34, 1.80, 0.18, size=10, bold=True, color=TEAL)
txt(s,
    "Instead of rebuilding the report after the field visit, LAIQ captures structured inspection data once, keeps evidence linked at source, and turns the same record into a report-ready package.",
    6.08, 5.60, 6.60, 0.34, size=9.2, color=DARK)

footer(s, "Tank Inspection · One-page solution summary")

prs.save(OUT)
print(f"✅  Saved → {OUT}")
