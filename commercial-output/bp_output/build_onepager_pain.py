"""
LAIQ Tank Inspection Pain One-Pager
Single-slide summary in the same visual language as the BP deck.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

BASE = os.path.dirname(os.path.abspath(__file__))
PARENT = os.path.dirname(BASE)
LOGO = os.path.join(PARENT, "laiq_logo.png")
OUT = os.path.join(BASE, "LAIQ_Tank_Pain_OnePager.pptx")

NAVY     = RGBColor(0x0A, 0x16, 0x28)
NAVY2    = RGBColor(0x12, 0x24, 0x40)
TEAL     = RGBColor(0x07, 0x5F, 0x5C)
TEAL_LT  = RGBColor(0xDC, 0xEF, 0xE9)
TEAL_MID = RGBColor(0x8A, 0xBB, 0xB6)
TEAL_DIM = RGBColor(0x4A, 0x80, 0x7C)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xF4, 0xF7, 0xF6)
LGRAY    = RGBColor(0xE2, 0xE8, 0xE6)
MID      = RGBColor(0x52, 0x62, 0x5E)
PALE     = RGBColor(0xA8, 0xB8, 0xB4)
AMBER    = RGBColor(0xE6, 0x7E, 0x22)
RED      = RGBColor(0xB0, 0x20, 0x20)
RED_LT   = RGBColor(0xFC, 0xEC, 0xEC)
DARK     = RGBColor(0x17, 0x20, 0x1E)

I = Inches
prs = Presentation()
prs.slide_width = I(13.33)
prs.slide_height = I(7.5)
BLANK = prs.slide_layouts[6]
FONT = "Arial"


def rect(s, l, t, w, h, fill=None, line=None, lw=0.6):
    sh = s.shapes.add_shape(1, I(l), I(t), I(w), I(h))
    sh.fill.solid() if fill else sh.fill.background()
    if fill:
        sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def txt(s, text, l, t, w, h, size=11, bold=False, color=DARK,
        align=PP_ALIGN.LEFT, wrap=True, italic=False):
    tb = s.shapes.add_textbox(I(l), I(t), I(w), I(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return tb


def bg(s, color=WHITE):
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = color


def add_logo(s, rgt=13.13, bot=7.44, h=0.26):
    if not os.path.exists(LOGO):
        return
    w = h * (930 / 430)
    s.shapes.add_picture(LOGO, I(rgt - w), I(bot - h), width=I(w), height=I(h))


def footer(s, label=""):
    rect(s, 0, 7.18, 13.33, 0.04, TEAL_MID)
    if label:
        txt(s, label, 0.18, 7.23, 6, 0.21, size=7.5, color=PALE)
    txt(s, "LAIQ  ·  Confidential", 7, 7.23, 6.13, 0.21,
        size=7.5, color=PALE, align=PP_ALIGN.RIGHT)
    add_logo(s)


def left_bar(s, color=TEAL, accent=TEAL_LT):
    rect(s, 0, 0, 0.32, 7.5, color)
    rect(s, 0.32, 0, 0.07, 7.5, accent)


def slide_kicker(s, kicker, title, tcolor=NAVY, kcolor=TEAL):
    txt(s, kicker, 0.55, 0.20, 12.6, 0.26, size=8.5, bold=True, color=kcolor)
    txt(s, title, 0.55, 0.44, 12.6, 0.76, size=24, bold=True, color=tcolor)
    rect(s, 0.55, 1.20, 12.60, 0.04, TEAL_MID)


s = prs.slides.add_slide(BLANK)
bg(s)
left_bar(s, TEAL)
slide_kicker(s, "PAIN POINTS",
             "Without LAIQ, Every Tank Report Is Rebuilt After the Field Visit", TEAL)

txt(s,
    "Oil & gas tank inspection is still highly manual. After the field visit, teams often rebuild the final report from notebooks, Excel sheets, photo folders, MFL/NDT attachments, and reviewer comments.",
    0.55, 1.32, 8.0, 0.40, size=10.5, color=MID)

# Left workflow panel
rect(s, 0.55, 1.82, 5.25, 3.72, OFFWHITE, TEAL_MID, 0.5)
txt(s, "Current broken handoff", 0.78, 1.98, 2.6, 0.22, size=10, bold=True, color=NAVY)
txt(s, "Field inspection → separate records → back-office re-keying → context rebuilding → report QA",
    0.78, 2.22, 4.6, 0.28, size=8.2, color=PALE, italic=True)

steps = [
    ("Field visit", TEAL),
    ("Notebook / photos /\nExcel", RED),
    ("Back-office\nre-keying", RED),
    ("Context rebuild", RED),
    ("Report QA", NAVY2),
]
sx = 0.86
for i, (label, color) in enumerate(steps):
    rect(s, sx, 2.66, 0.86, 0.64, color)
    txt(s, label, sx + 0.06, 2.78, 0.74, 0.34,
        size=8.2, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        txt(s, "▶", sx + 0.86, 2.84, 0.22, 0.22, size=17, bold=True, color=TEAL_MID, align=PP_ALIGN.CENTER)
    sx += 1.02

pain_cards = [
    ("Duplicate entry", "The same readings and notes are entered multiple times."),
    ("Weak evidence traceability", "Photos, findings, and location links are rebuilt manually."),
    ("Slow turnaround", "Report drafting starts after the site visit ends."),
    ("Low historical reuse", "Inspection knowledge remains trapped in static PDFs."),
]
for i, (title, body) in enumerate(pain_cards):
    top = 3.56 + i * 0.46
    rect(s, 0.86, top, 4.62, 0.36, WHITE, LGRAY, 0.3)
    rect(s, 0.86, top, 0.06, 0.36, RED)
    txt(s, title, 1.02, top + 0.05, 1.55, 0.16, size=8.8, bold=True, color=RED)
    txt(s, body, 2.48, top + 0.05, 2.82, 0.18, size=7.8, color=MID)

# Right quantified pain table
rect(s, 6.05, 1.82, 7.10, 3.72, WHITE, TEAL_MID, 0.5)
txt(s, "Pain area", 6.26, 2.00, 2.3, 0.20, size=10, bold=True, color=NAVY)
txt(s, "Baseline estimate", 10.20, 2.00, 2.0, 0.20, size=10, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
rect(s, 6.20, 2.26, 6.80, 0.03, LGRAY)

rows = [
    ("Post-site report assembly", "2–5 person-days/report"),
    ("Final report cycle after site", "7–10 days"),
    ("Mandatory structured points", "580+"),
    ("Total capture / verification burden", "1,000–2,000+ data points/report"),
    ("Checklist items", "~195 items"),
    ("UT readings", "~385 readings"),
    ("Evidence matching", "Manual photo–finding–location linking"),
    ("Reusable asset intelligence", "Low; trapped in static PDFs"),
]
for i, (label, value) in enumerate(rows):
    top = 2.40 + i * 0.36
    rect(s, 6.20, top, 6.80, 0.31, OFFWHITE if i % 2 == 0 else WHITE, LGRAY, 0.2)
    txt(s, label, 6.34, top + 0.05, 3.90, 0.16, size=8.6, color=DARK)
    txt(s, value, 10.38, top + 0.05, 2.42, 0.16, size=8.6, bold=True,
        color=RED if i < 4 else NAVY, align=PP_ALIGN.RIGHT)

# Bottom why it matters
rect(s, 0.55, 5.72, 12.60, 0.72, TEAL_LT, TEAL_MID, 0.5)
txt(s, "Why it matters", 0.82, 5.90, 1.6, 0.18, size=10, bold=True, color=TEAL)
txt(s,
    "A tank report is not just writing. It requires UT tables, checklist results, findings, severity, photos, MFL/NDT references, location records, and repair recommendations. Without structured capture, the same data is entered multiple times and evidence links are rebuilt manually.",
    2.18, 5.84, 10.55, 0.32, size=9.2, color=DARK)

rect(s, 0.55, 6.56, 12.60, 0.44, RED_LT, RED, 0.5)
txt(s, "LAIQ replaces manual post-site reconstruction with one structured capture path that feeds report generation directly.",
    0.88, 6.68, 11.94, 0.18, size=10.5, bold=True, color=RED, align=PP_ALIGN.CENTER)

footer(s, "Tank Inspection · One-page pain summary")

prs.save(OUT)
print(f"✅  Saved → {OUT}")
