from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os


OUT = "/Users/oscar/Documents/oscar-code/tank-inspection-coplilot/commercial-output/bp_output/LAIQ_Stage1_Budget_Oil_Gas_Inspection_Copilot.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_table(table, header_fill="DCEFE9"):
    table.style = "Table Grid"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, header_fill)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("LAIQ Stage 1 Budget\nOil & Gas Inspection Copilot")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(20)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Investor submission version")
r.italic = True
r.font.name = "Arial"
r.font.size = Pt(10)

doc.add_paragraph("")

doc.add_heading("1. Executive Summary", level=1)
for line in [
    "Stage 1 product: Oil & Gas Inspection Copilot.",
    "Delivery window: 6 months.",
    "Core engineering headcount ramps to 4 by Month 6.",
    "Full report generation is included by the end of Month 6.",
    "Estimated project budget: about SGD 250k, including modest contingency.",
]:
    p = doc.add_paragraph(style=None)
    p.style = doc.styles["Normal"]
    p.add_run("• ").bold = True
    p.add_run(line)

doc.add_paragraph("This plan is based on a tablet-first field product and a lean infrastructure model aligned to the current project scope.")

doc.add_heading("2. Stage 1 Product Scope", level=1)
scope_in = [
    "Inspection setup, scope, task board, and draft handling.",
    "Shell UT, roof UT, shell nozzle UT, roof nozzle UT.",
    "Inline findings with photos and evidence linking.",
    "Shell precise mapping.",
    "Bottom via MFL import / reference workflow.",
    "Review flow, export package, and full report generation by Month 6.",
    "Basic backend, authentication, storage, and pilot deployment readiness.",
]
doc.add_paragraph().add_run("Included deliverables").bold = True
for item in scope_in:
    p = doc.add_paragraph(style=None)
    p.add_run("• ").bold = True
    p.add_run(item)

doc.add_heading("3. Recommended Team", level=1)
team = doc.add_table(rows=1, cols=5)
team.rows[0].cells[0].text = "Role"
team.rows[0].cells[1].text = "Headcount"
team.rows[0].cells[2].text = "Start month"
team.rows[0].cells[3].text = "Responsibility"
team.rows[0].cells[4].text = "Delivery focus"
style_table(team)
rows = [
    ("Technical lead / senior full-stack", "1", "Month 1", "Architecture and delivery leadership", "System design, report pipeline, critical path"),
    ("Product engineer", "1", "Month 1", "Field workflow and UX", "Capture screens, offline flow, measurement UX"),
    ("Backend / platform engineer", "1", "Month 2", "Platform and report services", "API, persistence, files, report generation"),
    ("QA / automation engineer", "1", "Month 4", "Quality assurance and hardening", "Automation, device QA, pilot readiness"),
]
for row in rows:
    cells = team.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value

doc.add_paragraph("Product management and design are assumed to be covered internally. The only non-core delivery support budgeted here is part-time inspection subject matter expert (SME) / reviewer support.")

doc.add_paragraph().add_run("Core engineering ramp").bold = True
ramp = doc.add_table(rows=1, cols=3)
ramp.rows[0].cells[0].text = "Month"
ramp.rows[0].cells[1].text = "Core engineers active"
ramp.rows[0].cells[2].text = "Active roles"
style_table(ramp)
ramp_rows = [
    ("1", "2", "Technical lead, Product engineer"),
    ("2", "3", "Technical lead, Product engineer, Backend / platform engineer"),
    ("3", "3", "Technical lead, Product engineer, Backend / platform engineer"),
    ("4", "4", "Technical lead, Product engineer, Backend / platform engineer, QA / automation engineer"),
    ("5", "4", "Technical lead, Product engineer, Backend / platform engineer, QA / automation engineer"),
    ("6", "4", "Technical lead, Product engineer, Backend / platform engineer, QA / automation engineer"),
]
for row in ramp_rows:
    cells = ramp.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value

doc.add_heading("4. Six-Month Delivery Timeline", level=1)
doc.add_paragraph(
    "Stage 1 outcome at Month 6: the product must support end-to-end field capture and full report generation from captured inspection data."
)
timeline = doc.add_table(rows=1, cols=4)
timeline.rows[0].cells[0].text = "Month"
timeline.rows[0].cells[1].text = "Core engineers active"
timeline.rows[0].cells[2].text = "Focus"
timeline.rows[0].cells[3].text = "Output"
style_table(timeline)
months = [
    ("1", "2", "Foundations", "Data model, app shell, auth, storage, inspection lifecycle, report schema"),
    ("2", "3", "Core capture workflow", "Setup, scope, task board, shell UT, shell findings, shell mapping"),
    ("3", "3", "Remaining capture flows", "Roof UT, nozzle UT, MFL import, review flow, evidence management"),
    ("4", "4", "Report generation engine", "Template structure, table rendering, evidence linking, report assembly logic"),
    ("5", "4", "Full report generation + hardening", "End-to-end report output, offline behavior, QA automation, performance"),
    ("6", "4", "Pilot readiness", "Bug fixing, deployment, field feedback, report tuning, generated report acceptance"),
]
for row in months:
    cells = timeline.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value

doc.add_heading("5. Budget Summary", level=1)
summary = doc.add_table(rows=1, cols=2)
summary.rows[0].cells[0].text = "Item"
summary.rows[0].cells[1].text = "Budget"
style_table(summary)
summary_rows = [
    ("Estimated operating budget", "SGD 237.5k"),
    ("Contingency reserve", "SGD 12.5k"),
    ("Estimated total project budget", "SGD 250.0k"),
    ("Cloud services", "SGD 1.5k / month"),
    ("Coding / LLM API", "SGD 1.5k / month"),
    ("Full report generation", "Included by Month 6"),
]
for row in summary_rows:
    cells = summary.add_row().cells
    cells[0].text = row[0]
    cells[1].text = row[1]

doc.add_heading("6. Budget Breakdown", level=1)
bd = doc.add_table(rows=1, cols=3)
bd.rows[0].cells[0].text = "Budget line"
bd.rows[0].cells[1].text = "Budget"
bd.rows[0].cells[2].text = "Description"
style_table(bd)
breakdown = [
    ("Engineering team (ramp to 4 over 6 months)", "SGD 204.5k", "Core product build across app, backend, reporting, and QA"),
    ("Part-time inspection SME / reviewer", "SGD 10.0k", "Inspection-domain validation and report review"),
    ("Cloud service budget", "SGD 9.0k", "SGD 1.5k per month from Month 1"),
    ("Coding / LLM API budget", "SGD 9.0k", "SGD 1.5k per month from Month 1"),
    ("Devices, software, and test tooling", "SGD 5.0k", "Rugged-device testing, software, QA setup"),
    ("Contingency", "SGD 12.5k", "Lean delivery buffer and pilot-readiness reserve"),
]
for row in breakdown:
    cells = bd.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value

doc.add_heading("7. Monthly Expense Schedule", level=1)
doc.add_paragraph("All figures below are shown in SGD '000 and reflect estimated monthly operating burn. Cloud and coding / LLM API costs are budgeted from Month 1 onward. Contingency is shown separately.")
monthly = doc.add_table(rows=1, cols=7)
monthly.rows[0].cells[0].text = "Month"
monthly.rows[0].cells[1].text = "Engineering"
monthly.rows[0].cells[2].text = "Inspection SME"
monthly.rows[0].cells[3].text = "Cloud"
monthly.rows[0].cells[4].text = "Coding / LLM API"
monthly.rows[0].cells[5].text = "Devices / Tools"
monthly.rows[0].cells[6].text = "Total"
style_table(monthly)
monthly_rows = [
    ("1", "22.0", "1.0", "1.5", "1.5", "2.0", "28.0"),
    ("2", "32.0", "1.0", "1.5", "1.5", "1.5", "37.5"),
    ("3", "32.0", "1.5", "1.5", "1.5", "0.5", "37.0"),
    ("4", "39.5", "2.0", "1.5", "1.5", "0.5", "45.0"),
    ("5", "39.5", "2.0", "1.5", "1.5", "0.5", "45.0"),
    ("6", "39.5", "2.5", "1.5", "1.5", "0.0", "45.0"),
    ("Total", "204.5", "10.0", "9.0", "9.0", "5.0", "237.5"),
]
for row in monthly_rows:
    cells = monthly.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value

doc.add_paragraph(
    "Engineering burn is based on blended fully-loaded monthly costs for a Singapore / regional team: technical lead at about SGD 13k/month, product engineer at about SGD 9k/month, backend / platform engineer at about SGD 10k/month, and QA / automation engineer at about SGD 7.5k/month."
)
doc.add_paragraph("Contingency reserve: SGD 12.5k, held outside monthly operating burn and used only for delivery risk, pilot surprises, or final report-output hardening.")

doc.add_heading("8. Delivery Outcome at Month 6", level=1)
for item in [
    "Field capture workflow operational for Stage 1 oil & gas tank inspection scope.",
    "Shell, roof, nozzle, and MFL workflows complete.",
    "Evidence-linked inspection record available for review and export.",
    "Full report generation available from captured inspection data.",
    "Pilot-ready product release prepared for customer deployment.",
]:
    p = doc.add_paragraph(style=None)
    p.add_run("• ").bold = True
    p.add_run(item)

doc.add_heading("9. Budget Position", level=1)
doc.add_paragraph(
    "The Stage 1 budget is designed around a progressive staffing ramp to 4 engineers by Month 6 and includes full report generation within the 6-month delivery window."
)
doc.add_paragraph(
    "For investor submission and planning purposes, LAIQ should use an estimated Stage 1 project budget of about SGD 250k."
)

doc.save(OUT)
print(f"Saved {OUT}")
